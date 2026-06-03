import io
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.session import get_db
from app.dependencies import require_admin
from app.db.models.user import User
from app.db.models.connection import Connection
from app.db.models.guest import Guest
from app.db.models.employee import Employee
from app.db.models.audit_log import AuditLog
from app.db.models.browsing_log import BrowsingLog

router = APIRouter(prefix="/export", tags=["export"])


def _make_xlsx(headers: list, rows: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    # Header row
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4e73df")
        cell.alignment = Alignment(horizontal="center")
    # Data rows
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/connections")
async def export_connections(
    days: int = Query(30, ge=1, le=365),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
):
    since = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(days=days)
    result = await db.execute(
        select(Connection).where(Connection.connected_at >= since).order_by(Connection.connected_at.desc()).limit(5000)
    )
    rows = []
    for c in result.scalars():
        rows.append([
            c.id, c.hotspot_username, c.user_type, c.mac_address, c.ip_address,
            c.bytes_in, c.bytes_out, c.uptime_seconds,
            str(c.connected_at), str(c.disconnected_at or ""), c.disconnect_reason or "", c.is_active,
        ])
    headers = ["ID", "Username", "Type", "MAC", "IP", "Bytes In", "Bytes Out", "Uptime (s)",
               "Connected At", "Disconnected At", "Reason", "Active"]
    data = _make_xlsx(headers, rows)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=connections_{days}d.xlsx"},
    )


@router.get("/guests")
async def export_guests(
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
):
    result = await db.execute(select(Guest).order_by(Guest.created_at.desc()).limit(5000))
    rows = []
    for g in result.scalars():
        rows.append([
            g.id, g.full_name, g.email, g.mobile, g.status,
            str(g.created_at), str(g.approved_at or ""), str(g.access_expires_at or ""),
        ])
    headers = ["ID", "Name", "Email", "Mobile", "Status", "Created At", "Approved At", "Expires At"]
    data = _make_xlsx(headers, rows)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=guests.xlsx"},
    )


@router.get("/employees")
async def export_employees(
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
):
    result = await db.execute(select(Employee).order_by(Employee.created_at.desc()).limit(1000))
    rows = []
    for e in result.scalars():
        rows.append([
            e.id, e.full_name, e.email, e.hotspot_username, e.status,
            str(e.created_at),
        ])
    headers = ["ID", "Name", "Email", "Hotspot Username", "Status", "Created At"]
    data = _make_xlsx(headers, rows)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=employees.xlsx"},
    )


@router.get("/audit")
async def export_audit(
    days: int = Query(30, ge=1, le=365),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_admin),
):
    since = datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(days=days)
    result = await db.execute(
        select(AuditLog).where(AuditLog.created_at >= since).order_by(AuditLog.created_at.desc()).limit(10000)
    )
    rows = []
    for a in result.scalars():
        rows.append([a.id, a.username, a.action, a.resource_type or "", a.resource_id or "", a.ip_address or "", str(a.created_at)])
    headers = ["ID", "Username", "Action", "Resource Type", "Resource ID", "IP Address", "Timestamp"]
    data = _make_xlsx(headers, rows)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=audit_{days}d.xlsx"},
    )
